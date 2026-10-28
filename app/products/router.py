import csv
import io
import os
import uuid
from app.core.constants import UPLOAD_DIR
from fastapi import APIRouter, Depends, HTTPException, Response, UploadFile, status
from sqlalchemy.orm import Session

from app.cms.service import log_audit
from app.core.dependencies import get_db, require_admin, require_staff
from app.products.models import Category, Product
from app.products.schemas import (
    CategoryCreate,
    CategoryRead,
    CategoryUpdate,
    ProductCreate,
    ProductListItem,
    ProductRead,
    ProductUpdate,
)
from app.products.service import CategoryService, ProductService
from app.users.models import User
from pydantic import BaseModel

class PaginatedProducts(BaseModel):
    items: list[ProductListItem]
    total: int
    skip: int
    limit: int

router = APIRouter(prefix="/api", tags=["products"])

MAX_UPLOAD_SIZE = 5 * 1024 * 1024
ALLOWED_IMAGE_TYPES = {
    ".gif": ("image/gif", b"GIF8"),
    ".jpg": ("image/jpeg", b"\xff\xd8\xff"),
    ".jpeg": ("image/jpeg", b"\xff\xd8\xff"),
    ".png": ("image/png", b"\x89PNG\r\n\x1a\n"),
    ".webp": ("image/webp", b"RIFF"),
}


def product_with_stock(product: Product) -> ProductListItem:
    data = ProductRead.model_validate(product).model_dump()
    data["stock_quantity"] = product.inventory.quantity if product.inventory else 0
    return ProductListItem(**data)

# ── Categories
@router.get("/categories", response_model=list[CategoryRead])
def list_categories(db: Session = Depends(get_db)) -> list[Category]:
    return CategoryService(db).list_categories()


@router.post("/categories", response_model=CategoryRead, status_code=201, dependencies=[Depends(require_admin)])
def create_category(payload: CategoryCreate, db: Session = Depends(get_db)) -> Category:
    return CategoryService(db).create_category(payload)


@router.patch("/categories/{category_id}", response_model=CategoryRead, dependencies=[Depends(require_admin)])
def update_category(category_id: int, payload: CategoryUpdate, db: Session = Depends(get_db)) -> Category:
    return CategoryService(db).update_category(category_id, payload)


@router.delete("/categories/{category_id}", status_code=status.HTTP_204_NO_CONTENT, dependencies=[Depends(require_admin)])
def delete_category(category_id: int, db: Session = Depends(get_db)) -> Response:
    CategoryService(db).delete_category(category_id)
    return Response(status_code=status.HTTP_204_NO_CONTENT)


# ── Products
@router.get("/products", response_model=list[ProductListItem])
def list_products(
    search: str | None = None,
    category_id: int | None = None,
    min_price: float | None = None,
    max_price: float | None = None,
    in_stock: bool | None = None,
    skip: int = 0,
    limit: int = 50,
    db: Session = Depends(get_db),
) -> list[ProductListItem]:
    products = ProductService(db).list_products(
        search, category_id, min_price, max_price, in_stock, skip, limit
    )
    return [product_with_stock(p) for p in products]

@router.get("/products/paginated", response_model=PaginatedProducts)
def list_products_paginated(
    search: str | None = None,
    category_id: int | None = None,
    min_price: float | None = None,
    max_price: float | None = None,
    in_stock: bool | None = None,
    skip: int = 0,
    limit: int = 50,
    db: Session = Depends(get_db),
) -> PaginatedProducts:
    products = ProductService(db).list_products(search, category_id, min_price, max_price, in_stock, skip, limit)
    total = ProductService(db).count_products(search, category_id, min_price, max_price, in_stock)
    return PaginatedProducts(
        items=[product_with_stock(p) for p in products],
        total=total,
        skip=skip,
        limit=limit,
    )

@router.get("/products/{product_id}", response_model=ProductListItem)
def get_product(product_id: int, db: Session = Depends(get_db)) -> ProductListItem:
    return product_with_stock(ProductService(db).get_product(product_id))


@router.post("/products", response_model=ProductListItem, status_code=201)
def create_product(
    payload: ProductCreate,
    db: Session = Depends(get_db),
    admin: User = Depends(require_staff),
) -> ProductListItem:
    product = ProductService(db).create_product(payload)
    log_audit(db, admin.id, "create", "product", product.id, {"name": product.name})
    db.commit()
    return product_with_stock(product)


@router.patch("/products/{product_id}", response_model=ProductListItem)
def update_product(
    product_id: int,
    payload: ProductUpdate,
    db: Session = Depends(get_db),
    admin: User = Depends(require_staff),
) -> ProductListItem:
    product = ProductService(db).update_product(product_id, payload)
    log_audit(db, admin.id, "update", "product", product.id, {"name": product.name})
    db.commit()
    return product_with_stock(product)


@router.delete("/products/{product_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_product(
    product_id: int,
    db: Session = Depends(get_db),
    admin: User = Depends(require_staff),
) -> Response:
    ProductService(db).delete_product(product_id)
    log_audit(db, admin.id, "delete", "product", product_id)
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.post("/products/import", dependencies=[Depends(require_admin)])
async def import_products(file: UploadFile, db: Session = Depends(get_db), admin: User = Depends(require_admin)) -> dict:
    content = await file.read()
    name = (file.filename or "").lower()
    rows: list[dict] = []
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig")
        rows = list(csv.DictReader(io.StringIO(text)))
    elif name.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise HTTPException(status_code=400, detail="XLSX import requires openpyxl dependency") from exc
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        headers = [str(cell.value or "").strip() for cell in next(ws.iter_rows(max_row=1))]
        for line in ws.iter_rows(min_row=2, values_only=True):
            rows.append({headers[i]: line[i] for i in range(min(len(headers), len(line)))})
    else:
        raise HTTPException(status_code=400, detail="Upload CSV or XLSX")

    created = 0
    updated = 0
    errors = []
    category_service = CategoryService(db)
    product_service = ProductService(db)
    for idx, row in enumerate(rows, start=2):
        try:
            product_name = str(row.get("name") or row.get("Название") or row.get("ном") or "").strip()
            if not product_name:
                raise ValueError("name is required")
            category_name = str(row.get("category") or row.get("Категория") or row.get("категория") or "Imported").strip()
            category = next((c for c in category_service.list_categories() if c.name.lower() == category_name.lower()), None)
            if not category:
                category = category_service.create_category(CategoryCreate(name=category_name))
            price = row.get("price") or row.get("Цена") or row.get("нарх") or 1
            existing_slug = str(row.get("slug") or "").strip()
            existing = product_service.products.get_by_slug(existing_slug) if existing_slug else None
            payload = {
                "category_id": category.id,
                "name": product_name,
                "slug": existing_slug,
                "description": str(row.get("description") or row.get("Описание") or "—"),
                "price": price,
                "image_url": row.get("image_url") or row.get("image") or None,
                "unit": str(row.get("unit") or "шт"),
                "volume": str(row.get("volume") or "") or None,
                "initial_stock": int(row.get("stock") or row.get("остаток") or 0),
            }
            if existing:
                product_service.update_product(existing.id, ProductUpdate(**{k: v for k, v in payload.items() if k != "initial_stock"}))
                updated += 1
            else:
                product_service.create_product(ProductCreate(**payload))
                created += 1
        except Exception as exc:
            errors.append({"row": idx, "error": str(exc)})
    log_audit(db, admin.id, "import", "product", None, {"created": created, "updated": updated, "errors": len(errors)})
    db.commit()
    return {"created": created, "updated": updated, "errors": errors}


# ── Image Upload
@router.post("/upload-image", dependencies=[Depends(require_staff)])
async def upload_image(file: UploadFile) -> dict[str, str]:
    ext = os.path.splitext(file.filename or "img.png")[1].lower()
    expected = ALLOWED_IMAGE_TYPES.get(ext)
    if not expected:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Unsupported image type")
    expected_content_type, signature = expected
    if file.content_type != expected_content_type:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Image content type does not match extension")

    content = await file.read()
    if len(content) > MAX_UPLOAD_SIZE:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="Image is too large")
    if not content.startswith(signature):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid image file")
    if ext == ".webp" and content[8:12] != b"WEBP":
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid WebP file")

    filename = f"{uuid.uuid4().hex}{ext}"
    filepath = os.path.join(UPLOAD_DIR, filename)
    with open(filepath, "wb") as f:
        f.write(content)
    return {"url": f"/uploads/{filename}"}
